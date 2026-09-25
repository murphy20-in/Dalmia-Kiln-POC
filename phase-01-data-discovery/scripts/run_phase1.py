"""Run the complete Phase 1 pipeline and its validation.

  python3 run_phase1.py            # reuse parse cache if present
  python3 run_phase1.py --clean    # delete outputs/, reports/, cache/ first (clean execution)

Source integrity: SHA-256 of every source file is taken before and after the run
and compared; the run fails if anything changed.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
import pandas as pd

from common import SOURCE_ROOT, OUT_DIR, REPORT_DIR, CACHE_DIR, PHASE_DIR, sha256, jdump

STEPS = ["inventory_files.py", "inventory_workbooks.py", "profile_timestamps.py", "detect_gaps.py", "profile_columns.py",
         "profile_quality.py", "detect_outliers.py", "detect_flatlines.py", "build_inventories.py"]
REQUIRED = ["file_inventory.csv", "workbook_inventory.csv", "sheet_inventory.csv", "column_inventory.csv", "timestamp_inventory.csv",
            "data_quality_report.csv", "gap_analysis.csv", "outlier_profile.csv", "flatline_profile.csv", "unit_inventory.csv",
            "process_tag_inventory.csv", "equipment_data_matrix.csv", "event_data_inventory.csv", "alternative_fuel_inventory.csv",
            "phase1_data_readiness.csv"]


def snapshot() -> dict:
    return {str(p.relative_to(SOURCE_ROOT)): (sha256(p), p.stat().st_size, p.stat().st_mtime_ns)
            for p in sorted(SOURCE_ROOT.rglob("*")) if p.is_file()}


def run(script: str):
    t = time.time()
    print(f"[run] {script}")
    subprocess.run([sys.executable, script], cwd=Path(__file__).parent, check=True)
    print(f"      {time.time() - t:.1f}s")


def validate(before: dict) -> pd.DataFrame:
    V = []

    def chk(area, check, ok, evidence):
        V.append({"area": area, "check": check, "result": "PASS" if ok else "FAIL", "evidence": evidence})

    fi = pd.read_csv(OUT_DIR / "file_inventory.csv")
    src = fi[fi.scope == "SOURCE_DATA"]
    chk("File", "Every source file accounted for", set(src.relative_path) == set(before),
        f"{len(src)} inventoried vs {len(before)} on disk")
    hashes_match = all(before[r][0] == h for r, h in zip(src.relative_path, src.sha256))
    chk("File", "Inventory hashes equal pre-run snapshot", hashes_match, "sha256 per file")
    after = snapshot()
    chk("File", "No source modification (sha256, size, mtime before == after)", before == after,
        f"{sum(before[k] != after.get(k) for k in before)} changed, {len(set(after) - set(before))} added")
    for f in sorted(OUT_DIR.glob("*.csv")):
        try:
            n = len(pd.read_csv(f, low_memory=False))
            chk("File", f"Output opens: {f.name}", True, f"{n} rows")
        except Exception as e:
            chk("File", f"Output opens: {f.name}", False, str(e))
    missing = [r for r in REQUIRED if not (OUT_DIR / r).exists() and r != "phase1_data_readiness.csv"]
    chk("File", "All required outputs present (readiness written by generate_report)", not missing, jdump(missing))

    # schema: independent re-read of headers and dimensions with openpyxl
    si = pd.read_csv(OUT_DIR / "sheet_inventory.csv")
    ci = pd.read_csv(OUT_DIR / "column_inventory.csv", low_memory=False)
    cif = ci[ci.level == "FILE"]
    bad_dims, bad_hdr, bad_cols = [], [], []
    for r in si[si.sheet_kind == "PROCESS_LOG_TABLE"].itertuples():
        wb = openpyxl.load_workbook(SOURCE_ROOT / r.workbook, read_only=True)
        ws = wb[r.sheet_name]
        if ws.max_row != r.row_count_physical or ws.max_column != r.column_count_physical:
            bad_dims.append(r.workbook)
        hdr = [c for row in ws.iter_rows(min_row=1, max_row=int(r.data_start_row) - 1, values_only=True) for c in row]
        raw_text = " ".join(cif.loc[cif.source_file == r.workbook, "header_cells_raw"].astype(str)) + " " + str(r.title_text)
        for c in hdr:
            if isinstance(c, str) and c.strip() and c.strip() not in raw_text and c not in raw_text:
                bad_hdr.append(f"{r.workbook}:{c!r}")
        if (cif.source_file == r.workbook).sum() != r.column_count_physical:
            bad_cols.append(r.workbook)
        wb.close()
    chk("Schema", "Row/column counts re-verified with independent openpyxl read", not bad_dims, jdump(bad_dims))
    chk("Schema", "Column inventory has one row per physical column", not bad_cols, jdump(bad_cols))
    chk("Schema", "Original header text preserved verbatim (column headers + sheet title)", not bad_hdr, jdump(bad_hdr[:10]))
    chk("Schema", "Timestamp field identified in every process sheet", si.loc[si.sheet_kind == "PROCESS_LOG_TABLE", "data_start_row"].notna().all(),
        "column A, first timestamp row detected per sheet")
    ts = pd.read_csv(OUT_DIR / "timestamp_inventory.csv")
    tf = ts[ts.level == "FILE"]
    chk("Schema", "Timestamp parse success recorded", tf.parse_success_pct.notna().all(),
        f"min parse success {tf.parse_success_pct.min()}% across {len(tf)} files")
    chk("Schema", "Datatype assigned to every column", cif.datatype.notna().all(), f"{cif.datatype.value_counts().to_dict()}")
    # data
    dq = pd.read_csv(OUT_DIR / "data_quality_report.csv")
    ds = ci[ci.level == "DATASET"]
    chk("Data", "Missingness calculated for every dataset column", len(dq) == len(ds), f"{len(dq)} of {len(ds)}")
    gp = pd.read_csv(OUT_DIR / "gap_analysis.csv")
    chk("Data", "Gaps calculated for every file and dataset", len(gp) == len(tf) + ts.dataset.nunique(), f"{len(gp)} rows")
    chk("Data", "Duplicate timestamps checked", "duplicate_timestamp_rows" in gp.columns, f"{int(gp[gp.level == 'FILE'].duplicate_timestamp_rows.sum())} duplicate rows found")
    ol = pd.read_csv(OUT_DIR / "outlier_profile.csv")
    chk("Data", "Outliers profiled", len(ol) > 0, f"{len(ol)} rows, methods {sorted(ol.method.unique())}")
    fl = pd.read_csv(OUT_DIR / "flatline_profile.csv")
    chk("Data", "Flatlines checked for every dataset column", len(fl) == len(ds), f"{len(fl)} of {len(ds)}")
    un = pd.read_csv(OUT_DIR / "unit_inventory.csv")
    chk("Data", "Units documented for every dataset column", len(un) == len(ds) and un.original_unit.notna().all(), f"{len(un)} rows")
    return pd.DataFrame(V)


def main():
    if "--clean" in sys.argv:
        for d in (OUT_DIR, REPORT_DIR, CACHE_DIR):
            if d.exists():
                shutil.rmtree(d)
            d.mkdir(parents=True)
            if d != CACHE_DIR:
                (d / ".gitkeep").touch()  # tracked placeholder in the repo
        print("[clean] removed outputs/, reports/, cache/")
    t0 = time.time()
    before = snapshot()
    for s in STEPS:
        run(s)
    v = validate(before)
    v.to_csv(OUT_DIR / "pipeline_validation.csv", index=False)
    run("generate_report.py")
    run("export_data_manifest.py")
    for p in (REPORT_DIR / "PHASE_1_DATA_DISCOVERY_REPORT.md", REPORT_DIR / "PHASE_1_DATA_DISCOVERY_REPORT.html",
              OUT_DIR / "phase1_data_readiness.csv"):
        v.loc[len(v)] = ["File", f"Generated: {p.name}", "PASS" if p.exists() and p.stat().st_size > 0 else "FAIL",
                         f"{p.stat().st_size if p.exists() else 0} bytes"]
    final_ok = snapshot() == before
    v.loc[len(v)] = ["File", "No source modification after report generation", "PASS" if final_ok else "FAIL", "sha256/size/mtime"]
    v.to_csv(OUT_DIR / "pipeline_validation.csv", index=False)
    print(v.to_string(index=False, max_colwidth=90))
    print(f"[done] {time.time() - t0:.0f}s; {int((v.result == 'FAIL').sum())} validation failures")
    sys.exit(1 if (v.result == "FAIL").any() else 0)


if __name__ == "__main__":
    main()
